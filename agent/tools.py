"""
LangChain tools that the agent can use to interact with the system.

Action tools (click/type/...) operate on Windows 11 via PyAutoGUI. Mouse tools
accept an ``element_id`` (matching the numbered Set-of-Mark boxes / the UI-state
list) so the agent points at a real element instead of guessing coordinates;
raw (x, y) coordinates are still accepted as a fallback.

Also exposes Windows power tools (read on-screen text, create folders/files/Word
documents, run shell commands) and an explicit ``finish`` signal.
"""

from __future__ import annotations

import time
from typing import Any, Callable, Optional

from langchain.tools import tool

from actions.executor import ActionExecutor
from config import settings
from utils.logger import get_logger
from utils import win

log = get_logger(__name__)

# Global executor instance for tools to use
_executor = ActionExecutor()

# Set at runtime by the DigitalWorker (avoids circular imports).
_observe_fn: Optional[Callable[[], Any]] = None        # returns a Perception
_verify_fn: Optional[Callable[..., Any]] = None
_read_text_fn: Optional[Callable[..., str]] = None     # reads foreground window text
_current_goal: str = ""
_current_ui_state: str = ""
_current_elements: list = []
_current_image_scale: float = 1.0

# Completion signalling (set by the finish tool, read by the loop).
_task_finished: bool = False
_finish_summary: str = ""


def configure_tools(
    observe_fn: Callable[[], Any],
    verify_fn: Callable[..., Any],
    read_text_fn: Optional[Callable[..., str]] = None,
) -> None:
    """Wire up the worker callbacks used by tools."""
    global _observe_fn, _verify_fn, _read_text_fn
    _observe_fn = observe_fn
    _verify_fn = verify_fn
    _read_text_fn = read_text_fn
    log.info("Tools configured with observe/verify/read callbacks")


def set_current_context(
    goal: str,
    ui_state: str,
    elements: Optional[list] = None,
    image_scale: float = 1.0,
) -> None:
    """Update the current goal, UI-state text, element list, and image scale."""
    global _current_goal, _current_ui_state, _current_elements, _current_image_scale
    _current_goal = goal
    _current_ui_state = ui_state
    _current_elements = elements or []
    _current_image_scale = image_scale or 1.0


# --- completion helpers -----------------------------------------------------
def reset_finished() -> None:
    global _task_finished, _finish_summary
    _task_finished = False
    _finish_summary = ""


def is_finished() -> bool:
    return _task_finished


def get_finish_summary() -> str:
    return _finish_summary


# --- coordinate resolution --------------------------------------------------
def _resolve_xy(element_id: int, x: int, y: int) -> Optional[tuple[int, int, str]]:
    """Resolve a target point from an element ID or explicit coordinates.

    Returns (x, y, label) or None if it can't be resolved.
    """
    if element_id is not None and 0 <= element_id < len(_current_elements):
        el = _current_elements[element_id]
        return el.center_x, el.center_y, f'[{element_id}] "{el.text}"'
    if x is not None and x >= 0 and y is not None and y >= 0:
        # Raw coords come from the (possibly downscaled) image the model saw —
        # scale them back up to real screen pixels.
        scale = _current_image_scale or 1.0
        rx, ry = int(round(x / scale)), int(round(y / scale))
        return rx, ry, f"({rx}, {ry})"
    return None


def _verify_action(action_description: str, ui_state_before: str) -> str:
    """Optional post-action verification.

    Disabled by default: the agent re-observes a fresh screen each iteration and
    self-corrects. When ``VERIFY_AFTER_ACTION`` is on, captures the after-state
    and asks the LLM to judge; only undoes when ``VERIFY_AUTO_UNDO`` is also on.
    """
    if not settings.VERIFY_AFTER_ACTION or _verify_fn is None or _observe_fn is None:
        return f"OK: {action_description}"

    time.sleep(settings.VERIFICATION_DELAY)

    try:
        after = _observe_fn()
        ui_state_after = getattr(after, "text", str(after))
    except Exception as e:
        log.error("Failed to capture post-action UI state: %s", e)
        return f"OK: {action_description} (verification skipped: {e})"

    result = _verify_fn(
        action_description=action_description,
        ui_state_before=ui_state_before,
        ui_state_after=ui_state_after,
        overall_goal=_current_goal,
    )

    if result.verified:
        return f"VERIFIED: {action_description} — {result.explanation}"

    log.warning("Action verification FAILED: %s", result.explanation)
    undo_note = ""
    if settings.VERIFY_AUTO_UNDO:
        undo_result = _executor.undo()
        undo_note = f" Undo attempted ({undo_result})."

    hint = f" Suggested correction: {result.suggested_correction}" if result.suggested_correction else ""
    return (
        f"VERIFICATION FAILED: {action_description} — {result.explanation}.{undo_note}{hint} "
        f"Re-examine the fresh UI state and try a different approach."
    )


# ═══════════════════════════════════════════════════════════════════════════
# MOUSE ACTION TOOLS
# ═══════════════════════════════════════════════════════════════════════════

@tool
def click_element(element_id: int = -1, x: int = -1, y: int = -1, description: str = "") -> str:
    """Single-click a UI element.

    Args:
        element_id: ID of the target element from the UI-state list / numbered
            screenshot boxes. PREFERRED — use this whenever the element is listed.
        x: Fallback X coordinate (only if the element isn't in the list).
        y: Fallback Y coordinate.
        description: Short note about what is being clicked (for logging).
    """
    target = _resolve_xy(element_id, x, y)
    if target is None:
        return "Failed: provide a valid element_id (from the UI list) or x/y coordinates."
    tx, ty, label = target
    action_desc = f"Clicked {description or label} at ({tx}, {ty})"
    result = _executor.click(tx, ty)
    if result.startswith("Failed"):
        return result
    return _verify_action(action_desc, _current_ui_state)


@tool
def double_click_element(element_id: int = -1, x: int = -1, y: int = -1, description: str = "") -> str:
    """Double-click a UI element (open apps, files, folders).

    Prefer ``element_id``; fall back to ``x``/``y`` if needed.
    """
    target = _resolve_xy(element_id, x, y)
    if target is None:
        return "Failed: provide a valid element_id or x/y coordinates."
    tx, ty, label = target
    action_desc = f"Double-clicked {description or label} at ({tx}, {ty})"
    result = _executor.double_click(tx, ty)
    if result.startswith("Failed"):
        return result
    return _verify_action(action_desc, _current_ui_state)


@tool
def right_click_element(element_id: int = -1, x: int = -1, y: int = -1, description: str = "") -> str:
    """Right-click a UI element to open its Windows 11 context menu.

    Prefer ``element_id``; fall back to ``x``/``y`` if needed. If you need the
    full classic menu, click "Show more options" in the compact menu afterwards.
    """
    target = _resolve_xy(element_id, x, y)
    if target is None:
        return "Failed: provide a valid element_id or x/y coordinates."
    tx, ty, label = target
    action_desc = f"Right-clicked {description or label} at ({tx}, {ty})"
    result = _executor.right_click(tx, ty)
    if result.startswith("Failed"):
        return result
    return _verify_action(action_desc, _current_ui_state)


@tool
def scroll(direction: str, amount: int = 3, element_id: int = -1, x: int = -1, y: int = -1) -> str:
    """Scroll the mouse wheel up or down.

    Args:
        direction: "up" or "down".
        amount: Number of scroll clicks (larger = faster).
        element_id: Optional element to scroll over.
        x, y: Optional coordinates to scroll over (used if no element_id).
    """
    clicks = amount if direction.lower() == "up" else -amount
    target = _resolve_xy(element_id, x, y)
    pos_x, pos_y = (target[0], target[1]) if target else (None, None)
    action_desc = f"Scrolled {direction} {amount}"
    result = _executor.scroll(clicks, pos_x, pos_y)
    if result.startswith("Failed"):
        return result
    return _verify_action(action_desc, _current_ui_state)


@tool
def drag_element(
    start_element_id: int = -1,
    start_x: int = -1,
    start_y: int = -1,
    end_element_id: int = -1,
    end_x: int = -1,
    end_y: int = -1,
    description: str = "",
) -> str:
    """Drag from one point to another (move files, sliders, etc.).

    Specify the start and end either by element IDs or by coordinates.
    """
    start = _resolve_xy(start_element_id, start_x, start_y)
    end = _resolve_xy(end_element_id, end_x, end_y)
    if start is None or end is None:
        return "Failed: provide valid start/end element IDs or coordinates."
    action_desc = f"Dragged {description or start[2]} -> {end[2]}"
    result = _executor.drag(start[0], start[1], end[0], end[1])
    if result.startswith("Failed"):
        return result
    return _verify_action(action_desc, _current_ui_state)


# ═══════════════════════════════════════════════════════════════════════════
# KEYBOARD ACTION TOOLS
# ═══════════════════════════════════════════════════════════════════════════

@tool
def type_text(text: str, press_enter: bool = False) -> str:
    """Type a string via the keyboard into the currently focused field.

    Args:
        text: The text to type.
        press_enter: If True, press Enter right after typing. Use True for search
            bars, browser address bars, and any field you want to submit
            immediately (prevents autocomplete from corrupting the input).
    """
    action_desc = f'Typed "{text}"' + (" + Enter" if press_enter else "")
    result = _executor.type_text(text, press_enter=press_enter)
    if result.startswith("Failed"):
        return result
    return _verify_action(action_desc, _current_ui_state)


@tool
def press_key(key: str) -> str:
    """Press a single key (e.g., 'enter', 'tab', 'escape', 'down', 'f2')."""
    action_desc = f"Pressed '{key}'"
    result = _executor.press_key(key)
    if result.startswith("Failed"):
        return result
    return _verify_action(action_desc, _current_ui_state)


@tool
def press_hotkey(keys: str) -> str:
    """Press a key combination, e.g. 'win+s', 'ctrl+c', 'alt+f4', 'win+e'.

    Args:
        keys: Keys joined by '+', pressed simultaneously.
    """
    key_list = [k.strip() for k in keys.split("+")]
    action_desc = f"Pressed hotkey '{keys}'"
    result = _executor.hotkey(*key_list)
    if result.startswith("Failed"):
        return result
    return _verify_action(action_desc, _current_ui_state)


# ═══════════════════════════════════════════════════════════════════════════
# WINDOWS 11 POWER TOOLS
# ═══════════════════════════════════════════════════════════════════════════

@tool
def open_application(app_name: str) -> str:
    """Open a Windows 11 app via Start search (Win+S -> type -> Enter).

    Args:
        app_name: e.g. "Calculator", "Notepad", "File Explorer", "Edge",
            "Chrome", "Word", "Paint", "Terminal".
    """
    action_desc = f"Opened application '{app_name}'"
    if _executor.hotkey("win", "s").startswith("Failed"):
        return "Failed to open Windows Search."
    time.sleep(0.5)
    _executor.type_text(app_name)
    time.sleep(0.7)
    _executor.press_key("enter")
    time.sleep(1.0)
    return _verify_action(action_desc, _current_ui_state)


@tool
def read_screen_text(max_chars: int = 6000) -> str:
    """Read the visible text content of the CURRENT foreground window.

    Use this to "grab" information off the screen — e.g. news headlines in a
    browser, an article's text, search results, or a document's contents — so
    you can summarise or copy it elsewhere. Returns the extracted text.

    Args:
        max_chars: Maximum characters to return.
    """
    if _read_text_fn is None:
        return "read_screen_text is unavailable (no reader configured)."
    try:
        text = _read_text_fn(max_chars)
    except Exception as e:
        return f"Failed to read screen text: {e}"
    title = win.foreground_window_title() or "unknown window"
    if not text.strip():
        return f"(no readable text found in the foreground window: {title})"
    # Surface WHICH window was read so the model can tell if it read the wrong app.
    return f"--- TEXT FROM ACTIVE WINDOW: {title} ---\n{text}"


@tool
def create_folder(path: str) -> str:
    """Create a folder (and any parent folders).

    Accepts known-folder shortcuts, e.g. "desktop/News", "documents/Reports",
    or an absolute path like "C:\\Users\\me\\Desktop\\News".

    Args:
        path: Destination folder path.
    """
    try:
        target = win.resolve_path(path)
        target.mkdir(parents=True, exist_ok=True)
        return f"Created folder: {target}"
    except Exception as e:
        return f"Failed to create folder '{path}': {e}"


@tool
def write_text_file(path: str, content: str) -> str:
    """Write text to a .txt file (creating parent folders as needed).

    Accepts known-folder shortcuts, e.g. "desktop/News/summary.txt".

    Args:
        path: Destination file path.
        content: The full text to write.
    """
    try:
        target = win.resolve_path(path)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(content, encoding="utf-8")
        return f"Wrote {len(content)} chars to {target}"
    except Exception as e:
        return f"Failed to write file '{path}': {e}"


@tool
def create_document(path: str, content: str, title: str = "") -> str:
    """Create a document with the given content and save it.

    If ``path`` ends in ``.docx`` a real Word document is produced (requires
    python-docx); otherwise a ``.txt`` file is written. Accepts known-folder
    shortcuts, e.g. "desktop/News/Trending News Summary.docx".

    Args:
        path: Destination file path (.docx or .txt).
        content: Body text. Blank lines separate paragraphs.
        title: Optional heading placed at the top of the document.
    """
    try:
        target = win.resolve_path(path)
        target.parent.mkdir(parents=True, exist_ok=True)

        if target.suffix.lower() == ".docx":
            try:
                from docx import Document
            except Exception:
                # Fall back to a .txt sibling if python-docx isn't installed.
                target = target.with_suffix(".txt")
                body = (f"{title}\n{'=' * len(title)}\n\n" if title else "") + content
                target.write_text(body, encoding="utf-8")
                return f"python-docx missing; wrote text fallback to {target}"

            doc = Document()
            if title:
                doc.add_heading(title, level=0)
            for para in content.split("\n\n"):
                doc.add_paragraph(para.strip())
            doc.save(str(target))
            return f"Created Word document: {target}"

        body = (f"{title}\n{'=' * len(title)}\n\n" if title else "") + content
        target.write_text(body, encoding="utf-8")
        return f"Created document: {target}"
    except Exception as e:
        return f"Failed to create document '{path}': {e}"


@tool
def run_shell_command(command: str, timeout: int = 30) -> str:
    """Run a PowerShell command and return its output.

    Use for things that are far easier from the command line than the GUI
    (system info, listing files, network checks). Avoid destructive commands.

    Args:
        command: The PowerShell command to execute.
        timeout: Max seconds to wait (5–120).
    """
    safe_timeout = min(max(timeout, 5), 120)
    result = _executor.run_shell_command(command, timeout=safe_timeout)
    return f"Shell command: {command}\n{result}"


@tool
def wait(seconds: float = 1.0) -> str:
    """Pause briefly to let the UI catch up (e.g. an app finishing loading).

    Args:
        seconds: How long to wait (capped at 10s).
    """
    secs = min(max(seconds, 0.0), 10.0)
    time.sleep(secs)
    return f"Waited {secs:.1f}s"


@tool
def finish(summary: str) -> str:
    """Call this ONCE when the user's overall goal is fully accomplished.

    Args:
        summary: A short summary of what was accomplished.
    """
    global _task_finished, _finish_summary
    _task_finished = True
    _finish_summary = summary
    return f"DONE: {summary}"


def get_all_tools() -> list[Any]:
    """Return the tools available to the agent.

    Default (visible-GUI mode): mouse + keyboard + open_application + the
    read-only read_screen_text, so every artifact is produced by operating the
    real apps on screen. The silent file/folder/document/shell helpers are added
    only when ``ENABLE_POWER_TOOLS`` is set.
    """
    tools: list[Any] = [
        # Mouse
        click_element,
        double_click_element,
        right_click_element,
        scroll,
        drag_element,
        # Keyboard
        type_text,
        press_key,
        press_hotkey,
        # Visible app control + reading
        open_application,
        read_screen_text,
        wait,
        # Control
        finish,
    ]

    if settings.ENABLE_POWER_TOOLS:
        tools.extend([
            create_folder,
            write_text_file,
            create_document,
            run_shell_command,
        ])

    return tools
